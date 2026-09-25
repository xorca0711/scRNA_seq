"""Source-informed outcome reanalysis with animals and culture runs as units."""
import csv
import itertools
import json
import math
import statistics
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
import openpyxl
from a1_regulatory_fate import (BASE, describe, exact_permutation, pooled_fraction,
                                region_interaction, sha, sign_class, write_tsv)

CACHE = BASE / 'cache/regulatory_fate'
OUT = BASE / 'tables/regulatory_fate'


def ap1():
    path = CACHE / 'ap1_DC5.xlsx'
    w = openpyxl.load_workbook(path, read_only=True, data_only=True)
    fields, mice, groups, effects, omissions = [], [], [], [], []
    for sheet in ['wt SeV in situ', 'wt SeV de novo', 'mut SeV in situ', 'mut SeV de novo']:
        data = list(w[sheet].values)
        assert data[1][:4] == ('mouse', 'replicate', 'AT2 cells (GFP+)', 'HOPX+ AT2 cells')
        genotype = sheet.split()[0]
        region = 'in_situ' if 'in situ' in sheet else 'de_novo'
        for r, row in enumerate(data[2:11], 3):
            mouse, field, total, positive, percent = row
            assert mouse in [1, 2, 3] and field in [1, 2, 3]
            assert int(total) == total and int(positive) == positive
            fraction = pooled_fraction([positive], [total])
            assert math.isclose(100 * fraction, percent, rel_tol=1e-9)
            fields.append(dict(genotype=genotype, mouse=f'{genotype}_{mouse}', region=region,
                               field=field, positive=positive, total=total, percent=100 * fraction,
                               sheet=sheet, source_range=f'A{r}:E{r}'))
        for mouse in [1, 2, 3]:
            fs = [r for r in fields if r['mouse'] == f'{genotype}_{mouse}' and r['region'] == region]
            assert len(fs) == 3
            mice.append(dict(genotype=genotype, mouse=f'{genotype}_{mouse}', region=region,
                positive=sum(r['positive'] for r in fs), total=sum(r['total'] for r in fs),
                fields=len(fs), count_weighted_percent=100 * pooled_fraction([r['positive'] for r in fs], [r['total'] for r in fs]),
                equal_field_percent=statistics.mean(r['percent'] for r in fs)))
    for weighting in ['count_weighted_percent', 'equal_field_percent']:
        for genotype, region in itertools.product(['wt', 'mut'], ['in_situ', 'de_novo']):
            vals = [r[weighting] for r in mice if r['genotype'] == genotype and r['region'] == region]
            groups.append(dict(genotype=genotype, region=region, weighting=weighting, **describe(vals)))
        slopes = {}
        for genotype in ['wt', 'mut']:
            slopes[genotype] = []
            for mouse in [1, 2, 3]:
                rs = {r['region']: r[weighting] for r in mice if r['mouse'] == f'{genotype}_{mouse}'}
                slopes[genotype].append(rs['de_novo'] - rs['in_situ'])
        perm = exact_permutation(slopes['wt'], slopes['mut'])
        effects.append(dict(weighting=weighting, contrast='(mut_de_novo-mut_in_situ)-(wt_de_novo-wt_in_situ)', **perm))
        for wt, mut in itertools.product(range(3), repeat=2):
            effect = region_interaction([v for i, v in enumerate(slopes['wt']) if i != wt],
                                        [v for i, v in enumerate(slopes['mut']) if i != mut])
            omissions.append(dict(weighting=weighting, omitted_wt=f'wt_{wt+1}', omitted_mut=f'mut_{mut+1}',
                                  remaining_mice_per_genotype=2, interaction_pp=effect))
    return {'ap1_fields.tsv': fields, 'ap1_mice.tsv': mice, 'ap1_group_summary.tsv': groups,
            'ap1_region_interaction.tsv': effects, 'ap1_mouse_omissions.tsv': omissions}, [path]


def tsutsui():
    path = BASE / 'cache/followup_sources/tsutsui_source.xlsx'
    w = openpyxl.load_workbook(path, read_only=True, data_only=True)
    raw, summaries, contrasts = [], [], []

    def block(sheet, marker, rownum, label_col, first_col, last_col, mean_col, sd_col, panel, unit):
        row = list(w[sheet].values)[rownum - 1]
        label = row[label_col - 1]
        vals = []
        for col in range(first_col, last_col + 1):
            value = row[col - 1]
            if value in [None, '-']:
                continue
            assert isinstance(value, (int, float)) and not isinstance(value, bool)
            vals.append(value)
            raw.append(dict(panel=panel, marker=marker, condition=label,
                replicate_column=col-first_col+1, value=value, unit=unit, sheet=sheet,
                source_cell=f'{openpyxl.utils.get_column_letter(col)}{rownum}'))
        stats = describe(vals)
        assert math.isclose(stats['mean'], row[mean_col-1], rel_tol=1e-8, abs_tol=1e-10)
        assert math.isclose(stats['sd'], row[sd_col-1], rel_tol=1e-8, abs_tol=1e-10)
        summaries.append(dict(panel=panel, marker=marker, condition=label, unit=unit, **stats,
            sheet=sheet, source_range=f'{openpyxl.utils.get_column_letter(first_col)}{rownum}:{openpyxl.utils.get_column_letter(last_col)}{rownum}',
            unit_of_replication='author-declared independent culture experiment; one parental iPSC line'))

    s = list(w['Figure 6'].values)
    for header, left_panel, right_panel in [(4, '6c', '6d'), (10, '6c', '6d'), (16, '6c', '6d'),
                                           (25, '6g', '6h'), (31, '6g', '6h'), (37, '6g', '6h')]:
        for row in range(header+1, header+4):
            block('Figure 6', s[header-1][1], row, 2, 3, 9, 10, 11, left_panel, 'ratio_to_adult_lung_qPCR')
            block('Figure 6', s[header-1][13], row, 14, 15, 21, 22, 23, right_panel, 'ratio_to_adult_lung_qPCR')
    for row in range(47, 51):
        block('Figure 6', 'AGER_HiBiT', row, 2, 3, 5, 6, 7, '6i', 'luminescence_absolute_unit_unspecified')
    for row in range(5, 9):
        block('Figure 9', 'gel_contraction_inhibition', row, 2, 4, 6, 7, 8, '9d', 'normalized_percent_inhibition')
    s = list(w['Figure 9'].values)
    for header in [12, 16, 23, 27, 31, 35, 39]:
        for row in range(header+1, header+4):
            block('Figure 9', s[header-1][1], row, 2, 3, 6, 7, 8,
                  '9f' if header < 20 else '9g', 'relative_to_siCont')
    for r in summaries:
        if r['panel'] in ['6c', '6d', '6g', '6h'] and str(r['condition']).startswith('iATCs→'):
            ref = next(x for x in summaries if x['panel'] == r['panel'] and x['marker'] == r['marker'] and x['condition'] == 'iATCs_Day0')
            contrasts.append(dict(panel=r['panel'], marker=r['marker'], condition=r['condition'], comparator=ref['condition'],
                                   mean_ratio=r['mean']/ref['mean'], interpretation='medium-switch population differentiation capacity'))
        elif r['panel'] == '6i' and str(r['condition']).endswith('→iAT1s'):
            control = str(r['condition']).split('→')[0] + '→iAT2s'
            ref = next(x for x in summaries if x['panel'] == '6i' and x['condition'] == control)
            contrasts.append(dict(panel='6i', marker=r['marker'], condition=r['condition'], comparator=control,
                                   mean_ratio=r['mean']/ref['mean'], interpretation='AGER reporter induction, not mature AT1 function'))
        elif r['panel'] in ['9f', '9g'] and r['condition'] != 'siCont':
            contrasts.append(dict(panel=r['panel'], marker=r['marker'], condition=r['condition'], comparator='siCont',
                                   mean_ratio=r['mean'], interpretation='normalized expression; no re-test against artificial zero-variance controls'))

    genes_path = CACHE / 'tsutsui_data2.xlsx'
    gw = openpyxl.load_workbook(genes_path, read_only=True, data_only=True)
    rows = list(gw.active.values)
    genes = [{r[col] for r in rows[2:] if r[col] is not None} for col in range(4)]
    # The HNF1B-only column contains four duplicated symbols; retain source rows
    # and report unique-gene counts separately rather than counting them twice.
    assert [len(g) for g in genes] == [1339, 313, 325, 126]
    assert genes[0].isdisjoint(genes[1]) and genes[0].isdisjoint(genes[2]) and genes[1].isdisjoint(genes[2])
    assert genes[3] <= genes[2]
    memberships = []
    for rownum, row in enumerate(rows[2:], 3):
        for col, value in enumerate(row):
            if value is not None:
                memberships.append(dict(gene=value, supplied_set=rows[1][col].replace('\n', ' '),
                    source_cell=f'{openpyxl.utils.get_column_letter(col+1)}{rownum}',
                    evidence='author-derived motif-to-gene prediction and selected RNA overlap; no occupancy validation'))
    set_audit = []
    for col in range(4):
        values = [r[col] for r in rows[2:] if r[col] is not None]
        repeated = Counter(values)
        set_audit.append(dict(supplied_set=rows[1][col].replace('\n', ' '), source_rows=len(values),
            unique_genes=len(genes[col]), duplicated_symbols=','.join(sorted(g for g, n in repeated.items() if n > 1))))
    return {'tsutsui_source_values.tsv': raw, 'tsutsui_endpoint_summary.tsv': summaries,
            'tsutsui_descriptive_ratios.tsv': contrasts, 'tsutsui_regulatory_gene_sets.tsv': memberships,
            'tsutsui_gene_set_audit.tsv': set_audit}, [path, genes_path]


def tp53():
    path = CACHE / 'tp53_table1.xlsx'
    w = openpyxl.load_workbook(path, read_only=True, data_only=True)
    genes, counts = [], []
    focus = ['Itgb6', 'Krt8', 'Krt19', 'Cldn4', 'Ager', 'Sftpc', 'Cdkn1a', 'Krt7']
    for sheet in w:
        rows = list(sheet.values)
        seen, classification = set(), Counter()
        for rownum, row in enumerate(rows[1:], 2):
            gene = row[0]
            assert gene and gene not in seen
            seen.add(gene)
            if sheet.title == 'MDM2-KO Shared':
                a, b, qa, qb = row[2], row[8], row[6], row[12]
                assert abs(a) > 1 and abs(b) > 1 and qa < .05 and qb < .05
                category = sign_class(a, b)
            else:
                lfc, q = row[2], row[6]
                assert abs(lfc) > 1 and q < .05
                a, b = (lfc, '') if sheet.title.startswith('Ager') else ('', lfc)
                qa, qb = (q, '') if sheet.title.startswith('Ager') else ('', q)
                category = 'up' if lfc > 0 else 'down'
            classification[category] += 1
            genes.append(dict(gene=gene, supplied_set=sheet.title, ager_log2FC=a, sftpc_log2FC=b,
                              ager_author_q=qa, sftpc_author_q=qb, sign_class=category,
                              sheet=sheet.title, source_row=rownum))
        for category, n in classification.items():
            counts.append(dict(supplied_set=sheet.title, sign_class=category, genes=n, set_size=len(seen)))
    assert len(genes) == 3984 + 686 + 868
    all_names = [r['gene'] for r in genes]
    assert len(all_names) == len(set(all_names)), 'Unique and shared sets overlap unexpectedly'
    markers = []
    for gene in focus:
        rows = [r for r in genes if r['gene'] == gene]
        if rows:
            r = rows[0]
            markers.append(dict(gene=gene, supplied_set=r['supplied_set'], ager_log2FC=r['ager_log2FC'],
                sftpc_log2FC=r['sftpc_log2FC'], sign_class=r['sign_class'],
                limitation='unlisted contrast is unknown, not zero; no direct origin interaction tested'))
        else:
            markers.append(dict(gene=gene, supplied_set='not in selected lists', ager_log2FC='', sftpc_log2FC='',
                sign_class='unknown', limitation='absence from selected list is not evidence of no response'))
    return {'tp53_selected_gene_audit.tsv': genes, 'tp53_signed_set_summary.tsv': counts,
            'tp53_frozen_markers.tsv': markers}, [path]


def main():
    report = BASE / 'reports/regulatory_fate_analysis_run.json'
    assert not report.exists(), 'Refusing to overwrite analysis'
    outputs, inputs = {}, []
    for analysis in [ap1, tsutsui, tp53]:
        tables, sources = analysis()
        outputs.update(tables)
        inputs.extend(sources)
    assert all(not (OUT / name).exists() for name in outputs)
    for name, rows in outputs.items():
        write_tsv(OUT / name, rows)
    report.write_text(json.dumps(dict(utc=datetime.now(timezone.utc).isoformat(),
        code_sha256=sha(Path(__file__)), helper_sha256=sha(Path(__file__).with_name('a1_regulatory_fate.py')),
        contract_sha256=sha(BASE / 'config/regulatory_fate_analysis.json'),
        source_hashes={p.relative_to(BASE).as_posix(): sha(p) for p in inputs},
        output_hashes={str((OUT/name).relative_to(BASE)): sha(OUT/name) for name in outputs},
        rows={name: len(rows) for name, rows in outputs.items()},
        inference_limits='Selected-source extension; no raw-count fit, cell-as-replicate test or causal chromatin mediation'), indent=2) + '\n')
    print(json.dumps({name: len(rows) for name, rows in outputs.items()}))


if __name__ == '__main__':
    main()
